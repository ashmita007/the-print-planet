import os
import json
from datetime import datetime
from rest_framework import generics, viewsets, permissions, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.decorators import action
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework.authtoken.models import Token
from django.contrib.auth import authenticate, login
from django.contrib.auth.models import User
from django.db.models import Q, Count
from django.http import HttpResponse, Http404

from apps.catalog.views import StandardResultsSetPagination
from .models import BulkOrder, BulkOrderMessage, ConversationBackupArchive, UserCheckedProduct
from .serializers import (
    BulkOrderCreateSerializer,
    BulkOrderDetailSerializer,
    BulkOrderListSerializer,
    BulkOrderMessageSerializer,
    ConversationBackupArchiveSerializer,
    UserCheckedProductSerializer,
)


class BulkOrderSubmitView(APIView):
    """Submits a School or Corporate bulk order. Enforces or provides seamless user registration."""
    permission_classes = [permissions.AllowAny]
    throttle_scope = 'bulk_order'
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def post(self, request):
        user = request.user if request.user.is_authenticated else None
        token_key = None

        # If user is not logged in, registration or credentials are required
        if not user:
            email = request.data.get('email', '').strip().lower()
            password = request.data.get('password', '').strip()
            contact_person = request.data.get('contact_person', '').strip()
            organization_name = request.data.get('organization_name', '').strip()

            if not email:
                return Response(
                    {'error': 'Email is required for registration and order tracking.'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            # Check if user already exists
            existing_user = User.objects.filter(email__iexact=email).first()
            if existing_user:
                if not password:
                    return Response({
                        'error': 'An account with this email already exists. Please log in or enter your account password to submit.',
                        'requires_login': True
                    }, status=status.HTTP_401_UNAUTHORIZED)
                
                auth_user = authenticate(request, username=existing_user.username, password=password)
                if not auth_user:
                    return Response({'error': 'Incorrect password for existing account.'}, status=status.HTTP_401_UNAUTHORIZED)
                user = auth_user
            else:
                # Create a new user account
                if not password:
                    return Response(
                        {'error': 'Please set an account password so you can log in and track your bulk order conversation.'},
                        status=status.HTTP_400_BAD_REQUEST
                    )
                
                base_username = email.split('@')[0]
                base_username = ''.join(c for c in base_username if c.isalnum() or c in '._-') or 'buyer'
                username = base_username
                idx = 1
                while User.objects.filter(username=username).exists():
                    username = f"{base_username}{idx}"
                    idx += 1

                names = contact_person.split()
                first_name = names[0] if names else ''
                last_name = " ".join(names[1:]) if len(names) > 1 else ''

                user = User.objects.create_user(
                    username=username,
                    email=email,
                    password=password,
                    first_name=first_name,
                    last_name=last_name
                )

            token, _ = Token.objects.get_or_create(user=user)
            token_key = token.key
            login(request, user)

        # Process order creation
        serializer = BulkOrderCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        order = serializer.save(user=user)

        # Attach any products the user explicitly checked/explored
        checked_product_ids = request.data.get('checked_product_ids')
        if checked_product_ids:
            if isinstance(checked_product_ids, str):
                try:
                    checked_product_ids = json.loads(checked_product_ids)
                except Exception:
                    checked_product_ids = [int(x.strip()) for x in checked_product_ids.split(',') if x.strip().isdigit()]
            if isinstance(checked_product_ids, list):
                order.checked_products.set(checked_product_ids)

        # Set attachment url if uploaded
        if order.attachment:
            order.attachment_url = f"/api/v1/bulk-orders/my-orders/{order.order_number}/attachment/"
            order.save(update_fields=['attachment_url'])

        # Create initial conversation thread entry
        initial_msg_text = (
            f"Bulk Order Inquiry initiated for {order.organization_name} ({order.get_organization_type_display()}).\n"
            f"Quantity: {order.estimated_quantity} units | Category: {order.category_type}.\n\n"
            f"Requirements:\n{order.requirements}"
        )
        BulkOrderMessage.objects.create(
            order=order,
            sender=user,
            sender_type='client',
            sender_name=order.contact_person or user.get_full_name() or user.username,
            message=initial_msg_text,
            attachment_url=order.attachment_url or ''
        )

        # Automated welcome system message
        BulkOrderMessage.objects.create(
            order=order,
            sender=None,
            sender_type='system',
            sender_name='Print Planet Systems',
            message=(
                f"Thank you for your bulk inquiry #{order.order_number}! Our factory technical team at "
                f"Sunkadakatte, Bangalore is reviewing your specifications. A formal B2B quotation and "
                f"spec verification will be posted directly to this conversation thread within 2–4 business hours."
            )
        )

        return Response({
            'message': 'Bulk order inquiry submitted successfully!',
            'order_number': order.order_number,
            'token': token_key or getattr(getattr(user, 'auth_token', None), 'key', None),
            'order': BulkOrderDetailSerializer(order).data
        }, status=status.HTTP_201_CREATED)


class TrackCheckedProductView(APIView):
    """Records a product checked or browsed by an authenticated user for buyer insights."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, product_id):
        from apps.catalog.models import Product
        try:
            prod = Product.objects.get(id=product_id, is_active=True)
        except Product.DoesNotExist:
            return Response({'error': 'Product not found'}, status=status.HTTP_404_NOT_FOUND)

        item, created = UserCheckedProduct.objects.get_or_create(
            user=request.user,
            product=prod,
            defaults={'inquiry_count': 1}
        )
        if not created:
            item.inquiry_count += 1
            item.save(update_fields=['inquiry_count', 'viewed_at'])

        return Response({'success': True, 'checked_product_id': prod.id}, status=status.HTTP_200_OK)


class ClientCheckedProductsListView(generics.ListAPIView):
    """Lists all products previously checked by the current client."""
    serializer_class = UserCheckedProductSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = StandardResultsSetPagination

    def get_queryset(self):
        return UserCheckedProduct.objects.filter(
            user=self.request.user
        ).select_related('product', 'product__category').prefetch_related('product__images')


class ClientBulkOrdersListView(generics.ListAPIView):
    """Returns all bulk orders for the authenticated client with status counts."""
    serializer_class = BulkOrderListSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = StandardResultsSetPagination

    def get_queryset(self):
        user = self.request.user
        return BulkOrder.objects.filter(
            Q(user=user) | Q(email__iexact=user.email)
        ).select_related('product').prefetch_related('messages')

    def list(self, request, *args, **kwargs):
        response = super().list(request, *args, **kwargs)
        user = request.user
        qs = BulkOrder.objects.filter(Q(user=user) | Q(email__iexact=user.email))
        
        counts = {
            'total': qs.count(),
            'submitted': qs.filter(status='submitted').count(),
            'reviewing': qs.filter(status='reviewing').count(),
            'quoted': qs.filter(status='quoted').count(),
            'in_production': qs.filter(status__in=['proofing', 'in_production']).count(),
            'completed': qs.filter(status__in=['dispatched', 'completed']).count(),
        }
        response.data['summary_metrics'] = counts
        return response


class ClientBulkOrderDetailView(generics.RetrieveAPIView):
    """Retrieves full order details and message conversation for a specific bulk order."""
    serializer_class = BulkOrderDetailSerializer
    permission_classes = [permissions.IsAuthenticated]
    lookup_field = 'order_number'

    def get_queryset(self):
        user = self.request.user
        return BulkOrder.objects.filter(
            Q(user=user) | Q(email__iexact=user.email)
        ).select_related('product').prefetch_related('messages')


class ClientBulkOrderMessagePostView(APIView):
    """Allows client to send a message in the bulk order conversation."""
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def post(self, request, order_number):
        user = request.user
        order = BulkOrder.objects.filter(
            Q(order_number=order_number) & (Q(user=user) | Q(email__iexact=user.email))
        ).first()

        if not order:
            return Response({'error': 'Bulk order not found.'}, status=status.HTTP_404_NOT_FOUND)

        msg_text = request.data.get('message', '').strip()
        if not msg_text:
            return Response({'error': 'Message content cannot be empty.'}, status=status.HTTP_400_BAD_REQUEST)

        file_obj = request.FILES.get('attachment')
        attachment_url = ''
        if file_obj:
            # save attachment
            msg = BulkOrderMessage.objects.create(
                order=order,
                sender=user,
                sender_type='client',
                sender_name=order.contact_person or user.get_full_name() or user.username,
                message=msg_text,
                attachment=file_obj
            )
            attachment_url = f"/api/v1/bulk-orders/messages/{msg.id}/attachment/"
            msg.attachment_url = attachment_url
            msg.save(update_fields=['attachment_url'])
        else:
            msg = BulkOrderMessage.objects.create(
                order=order,
                sender=user,
                sender_type='client',
                sender_name=order.contact_person or user.get_full_name() or user.username,
                message=msg_text,
                attachment_url=request.data.get('attachment_url', '')
            )

        return Response(BulkOrderMessageSerializer(msg).data, status=status.HTTP_201_CREATED)


class OwnerBulkOrderViewSet(viewsets.ModelViewSet):
    """Owner management ViewSet for all bulk orders."""
    queryset = BulkOrder.objects.all().select_related('user', 'product').prefetch_related('messages')
    serializer_class = BulkOrderDetailSerializer
    permission_classes = [permissions.IsAuthenticated, permissions.IsAdminUser]
    pagination_class = StandardResultsSetPagination

    def get_queryset(self):
        qs = super().get_queryset()
        org_type = self.request.query_params.get('org_type')
        if org_type:
            qs = qs.filter(organization_type=org_type)
        
        status_param = self.request.query_params.get('status')
        if status_param:
            qs = qs.filter(status=status_param)

        q = self.request.query_params.get('search')
        if q:
            qs = qs.filter(
                Q(order_number__icontains=q) |
                Q(organization_name__icontains=q) |
                Q(contact_person__icontains=q) |
                Q(email__icontains=q) |
                Q(phone__icontains=q) |
                Q(city__icontains=q)
            )
        return qs

    def get_object(self):
        lookup = str(self.kwargs.get('pk', '')).strip()
        if lookup.isdigit():
            return super().get_object()
        obj = self.get_queryset().filter(order_number=lookup).first()
        if not obj:
            raise Http404("Bulk order not found.")
        self.check_object_permissions(self.request, obj)
        return obj

    @action(detail=True, methods=['post'], url_path='messages')
    def post_owner_reply(self, request, pk=None):
        order = self.get_object()
        msg_text = request.data.get('message', '').strip()
        if not msg_text:
            return Response({'error': 'Message cannot be empty.'}, status=status.HTTP_400_BAD_REQUEST)

        # Update order status or quote if provided
        new_status = request.data.get('status')
        if new_status and new_status in dict(BulkOrder.STATUS_CHOICES):
            order.status = new_status
        quoted_amt = request.data.get('quoted_amount')
        if quoted_amt is not None and str(quoted_amt).strip() != '':
            try:
                order.quoted_amount = float(quoted_amt)
            except ValueError:
                pass
        tracking_num = request.data.get('tracking_consignment_number')
        if tracking_num:
            order.tracking_consignment_number = tracking_num
        courier = request.data.get('tracking_courier_partner')
        if courier:
            order.tracking_courier_partner = courier
        order.save()

        msg = BulkOrderMessage.objects.create(
            order=order,
            sender=request.user,
            sender_type='owner',
            sender_name="The Print Planet Support",
            message=msg_text,
            attachment_url=request.data.get('attachment_url', '')
        )

        return Response({
            'message': BulkOrderMessageSerializer(msg).data,
            'order': BulkOrderDetailSerializer(order).data
        }, status=status.HTTP_201_CREATED)


class OwnerConversationBackupDownloadView(APIView):
    """Separate dedicated backup generator for all user conversations and bulk orders in Excel format."""
    permission_classes = [permissions.IsAuthenticated, permissions.IsAdminUser]

    def get(self, request):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

        # ----------------------------------------------------
        # Sheet 1: Bulk Orders Overview
        # ----------------------------------------------------
        ws1 = wb.active
        ws1.title = "Bulk Orders Summary"
        h1 = [
            'Order #', 'Organization Name', 'Org Type', 'Contact Person', 'Email',
            'Phone', 'City', 'Category', 'Quantity', 'Status', 'Quoted Amount (INR)',
            'Tracking Courier', 'Tracking AWB', 'Messages Count', 'Created Date'
        ]
        ws1.append(h1)
        for c_idx in range(1, len(h1) + 1):
            cell = ws1.cell(row=1, column=c_idx)
            cell.fill = PatternFill(start_color='17365D', end_color='17365D', fill_type='solid')
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        orders = BulkOrder.objects.select_related('user', 'product').prefetch_related('messages').all()
        for o in orders:
            ws1.append([
                o.order_number,
                o.organization_name,
                o.get_organization_type_display(),
                o.contact_person,
                o.email,
                o.phone,
                o.city,
                o.category_type,
                o.estimated_quantity,
                o.get_status_display(),
                float(o.quoted_amount) if o.quoted_amount else '',
                o.tracking_courier_partner,
                o.tracking_consignment_number,
                o.messages.count(),
                o.created_at.strftime('%Y-%m-%d %H:%M:%S')
            ])

        for col in ws1.columns:
            ws1.column_dimensions[get_column_letter(col[0].column)].width = 24

        # ----------------------------------------------------
        # Sheet 2: All Conversation Messages
        # ----------------------------------------------------
        ws2 = wb.create_sheet(title="User Conversations Backup")
        h2 = ['Message ID', 'Order #', 'Organization', 'Sender Type', 'Sender Name', 'Message Content', 'Timestamp']
        ws2.append(h2)
        for c_idx in range(1, len(h2) + 1):
            cell = ws2.cell(row=1, column=c_idx)
            cell.fill = PatternFill(start_color='E87832', end_color='E87832', fill_type='solid')
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        msgs = BulkOrderMessage.objects.select_related('order').all().order_by('id')
        for m in msgs:
            ws2.append([
                m.id,
                m.order.order_number,
                m.order.organization_name,
                m.get_sender_type_display(),
                m.sender_name,
                m.message,
                m.created_at.strftime('%Y-%m-%d %H:%M:%S')
            ])

        for col in ws2.columns:
            ws2.column_dimensions[get_column_letter(col[0].column)].width = 26

        buffer = openpyxl.writer.excel.save_virtual_workbook(wb)
        timestamp = datetime.now().strftime('%Y%m%d-%H%M')
        filename = f"the-print-planet-bulk-orders-and-conversations-{timestamp}.xlsx"

        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class OwnerConversationBackupCreateSnapshotView(APIView):
    """Creates a persistent conversation archive record in database for safety audit."""
    permission_classes = [permissions.IsAuthenticated, permissions.IsAdminUser]

    def post(self, request):
        timestamp = datetime.now().strftime('%Y%m%d-%H%M%S')
        archive_code = f"ARCH-CONV-{timestamp}"

        orders_data = []
        for o in BulkOrder.objects.prefetch_related('messages').all():
            messages_list = [
                {
                    'id': m.id,
                    'sender_type': m.sender_type,
                    'sender_name': m.sender_name,
                    'message': m.message,
                    'created_at': m.created_at.isoformat()
                }
                for m in o.messages.all()
            ]
            orders_data.append({
                'order_number': o.order_number,
                'organization_name': o.organization_name,
                'organization_type': o.organization_type,
                'contact_person': o.contact_person,
                'email': o.email,
                'phone': o.phone,
                'status': o.status,
                'quantity': o.estimated_quantity,
                'requirements': o.requirements,
                'created_at': o.created_at.isoformat(),
                'messages': messages_list
            })

        total_orders = len(orders_data)
        total_messages = sum(len(o['messages']) for o in orders_data)

        archive = ConversationBackupArchive.objects.create(
            archive_code=archive_code,
            total_orders=total_orders,
            total_messages=total_messages,
            backup_data={'orders': orders_data},
            created_by=request.user,
            notes=request.data.get('notes', 'Manual owner conversation backup snapshot')
        )

        return Response({
            'message': 'Conversation snapshot backed up successfully',
            'archive_code': archive.archive_code,
            'total_orders': archive.total_orders,
            'total_messages': archive.total_messages,
            'created_at': archive.created_at
        }, status=status.HTTP_201_CREATED)
